"""File readers: large-file-safe text chunks, CSV and XLSX previews."""

from __future__ import annotations

import csv
import io
import sys

KB = 1024
MB = 1024 * KB
CONTENT_MAX = 2 * MB  # /api/content limit
SMALL_TEXT = 256 * KB  # whole file in one segment at or below this
HEAD_TAIL = 128 * KB
CHUNK_MAX = 1 * MB  # hard server-side cap on any requested length
LINE_WINDOW = 256 * KB
CSV_MAX_ROWS = 5000
XLSX_MAX_ROWS = 1000
_BLOCK = 1 * MB

csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


# --- text chunks -------------------------------------------------------------
def _utf8_forward(data: bytes, i: int) -> int:
    while i < len(data) and (data[i] & 0xC0) == 0x80:
        i += 1
    return i


def _utf8_back(data: bytes, i: int) -> int:
    while i > 0 and i < len(data) and (data[i] & 0xC0) == 0x80:
        i -= 1
    return i


def count_newlines(f, end: int) -> int:
    f.seek(0)
    n = 0
    pos = 0
    while pos < end:
        block = f.read(min(_BLOCK, end - pos))
        if not block:
            break
        n += block.count(b"\n")
        pos += len(block)
    return n


def offset_of_line(f, line: int, size: int) -> int:
    """Byte offset of the start of 1-based line `line` (clamped to the last line)."""
    if line <= 1:
        return 0
    f.seek(0)
    need = line - 1
    pos = 0
    last_start = 0
    while pos < size:
        block = f.read(_BLOCK)
        if not block:
            break
        i = -1
        while True:
            i = block.find(b"\n", i + 1)
            if i == -1:
                break
            need -= 1
            start = pos + i + 1
            if start < size:
                last_start = start
            if need == 0:
                return min(start, size)
        pos += len(block)
    return last_start


def read_segment(f, size: int, start: int, length: int) -> dict:
    """Read [start, start+length) snapped to line boundaries, never splitting UTF-8."""
    start = max(0, min(int(start), size))
    length = max(0, min(int(length), CHUNK_MAX))
    end = min(size, start + length)
    # read one extra byte before start so we know whether start is already at a line start
    lead = 1 if start > 0 else 0
    f.seek(start - lead)
    data = f.read(end - start + lead)
    base = start - lead
    s = lead
    e = len(data)
    if start > 0 and data[0:1] != b"\n":
        nl = data.find(b"\n", s, e)
        if nl != -1:
            s = nl + 1
        else:
            s = _utf8_forward(data, s)  # one giant line: fall back to a UTF-8 boundary
    if end < size:
        nl = data.rfind(b"\n", s, e)
        if nl != -1:
            e = nl + 1
        else:
            e = _utf8_back(data, e)
            if e < s:
                e = s
    seg_start, seg_end = base + s, base + e
    return {
        "start": seg_start,
        "end": seg_end,
        "start_line": count_newlines(f, seg_start) + 1,
        "text": data[s:e].decode("utf-8", errors="replace"),
    }


def _gaps(size: int, segments: list[dict]) -> list[dict]:
    gaps, pos = [], 0
    for seg in sorted(segments, key=lambda s: s["start"]):
        if seg["start"] > pos:
            gaps.append({"start": pos, "end": seg["start"]})
        pos = max(pos, seg["end"])
    if pos < size:
        gaps.append({"start": pos, "end": size})
    return gaps


def text_chunks(real: str, size: int, start: int | None = None, length: int | None = None,
                line: int | None = None) -> dict:
    with open(real, "rb") as f:
        if line is not None:
            off = offset_of_line(f, max(1, line), size)
            seg = read_segment(f, size, max(0, off - LINE_WINDOW // 2), LINE_WINDOW)
            if not (seg["start"] <= off < seg["end"] or off >= size):  # extremely long line: start at it
                seg = read_segment(f, size, off, LINE_WINDOW)
            segments = [seg]
        elif start is not None:
            segments = [read_segment(f, size, start, HEAD_TAIL if length is None else length)]
        elif size <= SMALL_TEXT:
            segments = [read_segment(f, size, 0, size)]
        else:
            head = read_segment(f, size, 0, HEAD_TAIL)
            tail = read_segment(f, size, size - HEAD_TAIL, HEAD_TAIL)
            segments = [head, tail] if tail["start"] > head["end"] else [read_segment(f, size, 0, CHUNK_MAX)]
    return {"size": size, "segments": segments, "gaps": _gaps(size, segments)}


# --- CSV -----------------------------------------------------------------------
def read_csv(real: str, ext: str) -> dict:
    with open(real, "r", encoding="utf-8", errors="replace", newline="") as f:
        sample = f.read(64 * KB)
        f.seek(0)
        if ext == "tsv":
            dialect: type[csv.Dialect] | csv.Dialect = csv.excel_tab
        else:
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
        reader = csv.reader(f, dialect)
        header = next(reader, [])
        rows, total = [], 0
        for row in reader:
            total += 1
            if total <= CSV_MAX_ROWS:
                rows.append(row)
    return {"header": header, "rows": rows, "total_rows": total, "truncated": total > CSV_MAX_ROWS}


# --- XLSX ----------------------------------------------------------------------
def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_xlsx(real: str) -> dict:
    from openpyxl import load_workbook

    with open(real, "rb") as fh:
        data = io.BytesIO(fh.read())
    wb = load_workbook(data, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows, total = [], 0
            for row in ws.iter_rows(values_only=True):
                total += 1
                if total <= XLSX_MAX_ROWS:
                    rows.append([_cell(v) for v in row])
            sheets.append({"name": ws.title, "rows": rows, "total_rows": total,
                           "truncated": total > XLSX_MAX_ROWS})
        return {"sheets": sheets}
    finally:
        wb.close()
